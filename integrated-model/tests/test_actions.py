import json,os,tempfile,unittest,zipfile,hashlib,math
from pathlib import Path
from unittest.mock import patch
import openpyxl
from cryptography.fernet import Fernet,InvalidToken
from railmodel.batch import execute,geometry_inventory
from railmodel.bridge import encrypt_bytes,decrypt_bytes,encrypt_output,decrypt_output,materialize_request
from railmodel.fixture import synthetic_data
from railmodel.engine import Engine

class ActionsTests(unittest.TestCase):
 def test_synthetic_six_run_export_and_formulas(self):
  with tempfile.TemporaryDirectory()as t:
   out=Path(t)/'runs';m=execute(None,out,{'year':2045,'case':'Base'},True,synthetic=True)
   self.assertEqual(m['run_count'],6);self.assertTrue(m['synthetic'])
   w=openpyxl.load_workbook(out/'Sabah_Rail_Audit.xlsx',data_only=True)
   self.assertEqual(len(w.sheetnames),20)
   self.assertTrue(all(w['QA'].cell(i,2).value<1e-7 for i in [4,5,6]))
   self.assertEqual(w['Passenger Calc']['B5'].value,w['Controls']['B4'].value*w['Controls']['B5'].value)
   self.assertFalse([(s.title,c.coordinate)for s in w for row in s for c in row if c.data_type=='e'])
   f=openpyxl.load_workbook(out/'Sabah_Rail_Audit.xlsx',data_only=False)
   self.assertEqual(f['Passenger Calc']['B5'].value,'=B4*Controls!B5')
   runs=json.loads((out/'runs.json').read_text());self.assertEqual({r['validation_status']for r in runs},{'SYNTHETIC_TEST_ONLY'})
   for item in m['files']:self.assertEqual(item['sha256'],hashlib.sha256((out/item['name']).read_bytes()).hexdigest())
 def test_approval_gate_never_infers_latest(self):
  with tempfile.TemporaryDirectory()as t:
   path=Path(t)/'data.json';data=synthetic_data();data['governance_status']='Draft evidence';path.write_text(json.dumps(data))
   with self.assertRaisesRegex(ValueError,'Approved numerical'):execute(path,Path(t)/'runs')
 def test_encryption_roundtrip_and_wrong_key(self):
  key=Fernet.generate_key();other=Fernet.generate_key();value=b'CONFIDENTIAL test bytes';encrypted=encrypt_bytes(value,key)
  self.assertNotIn(value,encrypted);self.assertEqual(decrypt_bytes(encrypted,key),value)
  with self.assertRaises(InvalidToken):decrypt_bytes(encrypted,other)
 def test_encrypted_output_roundtrip(self):
  with tempfile.TemporaryDirectory()as t:
   root=Path(t);out=root/'plain';out.mkdir();(out/'proof.txt').write_text('synthetic proof');key=Fernet.generate_key()
   encrypt_output(out,root/'out.enc',key);decrypt_output(root/'out.enc',root/'restored',key)
   self.assertEqual((root/'restored/proof.txt').read_text(),'synthetic proof')
 def test_unsafe_archive_rejected(self):
  import io
  b=io.BytesIO()
  with zipfile.ZipFile(b,'w')as z:z.writestr('../escape','bad')
  with tempfile.TemporaryDirectory()as t:
   root=Path(t);key=Fernet.generate_key();(root/'in.enc').write_bytes(encrypt_bytes(b.getvalue(),key))
   with self.assertRaisesRegex(ValueError,'Unsafe archive'):decrypt_output(root/'in.enc',root/'out',key)
 def test_ciphertext_tamper_rejected(self):
  with tempfile.TemporaryDirectory()as t:
   root=Path(t);(root/'input.enc').write_bytes(b'tamper');(root/'request.json').write_text(json.dumps({'schema':'sabah-work-bridge/1','ciphertext_sha256':'incorrect'}))
   with self.assertRaisesRegex(ValueError,'checksum'):materialize_request(root/'request.json',root/'data.json',Fernet.generate_key())
 def test_missing_secret_fails_closed(self):
  with patch.dict(os.environ,{'SABAH_RAIL_BRIDGE_KEY':''}):
   with self.assertRaisesRegex(ValueError,'locked'):encrypt_bytes(b'data')
 def test_wgs84_inventory_and_invalid_coordinates(self):
  a={'type':'FeatureCollection','features':[{'type':'Feature','properties':{'name':'Synthetic'},'geometry':{'type':'LineString','coordinates':[[110,5],[111,6]]}}]}
  r=geometry_inventory(a);self.assertEqual(r[0]['coordinate_count'],2);self.assertEqual(r[0]['xmin'],110)
  a['features'][0]['geometry']['coordinates']=[[190,6]]
  with self.assertRaises(ValueError):geometry_inventory(a)
 def test_invalid_headway_and_variant(self):
  with tempfile.TemporaryDirectory()as t:
   path=Path(t)/'data.json';path.write_text(json.dumps(synthetic_data()));e=Engine(path)
   with self.assertRaises(ValueError):e.run({'overrides':{'headway_min':0}})
   with self.assertRaises(ValueError):e.run({'passenger_variant':'tn3_high_only','case':'High','option':1})

if __name__=='__main__':unittest.main()
